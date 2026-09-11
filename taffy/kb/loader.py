"""文档解析：把 knowledge/ 里的各种格式统一转成纯文本。

对外只暴露 supported() 和 load()。load() 返回 [(文本段, 页码或 None), ...]，
PDF 按页返回所以带页码，其它格式页码是 None。解析失败会直接抛异常，
由上层记录下来跳过这个文件，不影响其它文档。
"""
import os
from html.parser import HTMLParser

# 纯文本类，直接读文件
_TEXT_EXTS = (".txt", ".md", ".markdown", ".log")

# 这些标签里的内容是代码或样式，不要
_SKIP_TAGS = {"script", "style", "noscript"}


class _HTMLExtractor(HTMLParser):
    """剥掉标签只留正文，script / style 整段忽略。"""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._parts = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in _SKIP_TAGS:
            self._skip_depth += 1

    def handle_endtag(self, tag):
        if tag in _SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data):
        if not self._skip_depth and data.strip():
            self._parts.append(data.strip())

    @property
    def text(self):
        return "\n".join(self._parts)


def _read_text(path):
    """按 utf-8 -> gbk 试编码，都不行就用替换字符硬解，保证不炸。"""
    for encoding in ("utf-8", "gbk"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _load_plain(path):
    return [(_read_text(path), None)]


def _load_html(path):
    parser = _HTMLExtractor()
    parser.feed(_read_text(path))
    return [(parser.text, None)]


def _load_pdf(path):
    """按页抽取文字，顺带把页码带出去，方便回答里标出处。"""
    from pypdf import PdfReader

    pages = []
    for number, page in enumerate(PdfReader(path).pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            # 单页坏掉就跳过这页，不要整个文件报废
            continue
        if text.strip():
            pages.append((text, number))
    return pages


def _load_docx(path):
    """段落 + 表格一起抽，表格用制表符拼成行。"""
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return [("\n".join(parts), None)]


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".html": _load_html,
    ".htm": _load_html,
}
for _ext in _TEXT_EXTS:
    _LOADERS[_ext] = _load_plain


def supported(path):
    """这个后缀的知识库支持吗"""
    return os.path.splitext(path)[1].lower() in _LOADERS


def load(path):
    """解析一个文档，返回 [(文本段, 页码或 None)]；不支持的格式返回空列表。"""
    loader = _LOADERS.get(os.path.splitext(path)[1].lower())
    if loader is None:
        return []
    return [(text, page) for text, page in loader(path) if text and text.strip()]
